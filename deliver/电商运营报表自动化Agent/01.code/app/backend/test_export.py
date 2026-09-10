"""S5 导出接口测试：Markdown 原样导出 + Word(docx) 结构化渲染。

运行：从项目根目录 `python -m pytest app/backend/test_export.py -q`
生成采用 ?sync=true，保证确定性；测试库由 conftest 隔离（app_test.db）。
"""
from io import BytesIO

import pytest
from docx import Document as DocxDocument
from fastapi.testclient import TestClient

from app.backend.db import SessionLocal, init_db
from app.backend.main import app
from app.backend.models import Report, User

init_db()

SAMPLE_INSTR = "2010年11月22日到28日的日报"  # 促销周样例：GMV £296,087.91


@pytest.fixture(autouse=True)
def _reset():
    db = SessionLocal()
    try:
        db.query(Report).delete()
        db.query(User).delete()
        db.commit()
    finally:
        db.close()
    yield


@pytest.fixture
def client():
    return TestClient(app)


@pytest.fixture
def auth_headers(client):
    r = client.post("/api/auth/register", json={"username": "alice", "password": "secret123"})
    assert r.status_code == 200, r.text
    return {"Authorization": f"Bearer {r.json()['access_token']}"}


def _gen_sync(client, headers, instruction=SAMPLE_INSTR) -> int:
    r = client.post(
        "/api/reports/generate",
        params={"sync": True},
        json={"instruction": instruction},
        headers=headers,
    )
    assert r.status_code == 200, r.text
    assert r.json()["status"] == "drafted"
    return r.json()["id"]


def _docx_text(blob: bytes) -> str:
    d = DocxDocument(BytesIO(blob))
    parts = [p.text for p in d.paragraphs]
    for t in d.tables:
        for row in t.rows:
            for cell in row.cells:
                parts.append(cell.text)
    return "\n".join(parts)


def test_export_requires_auth(client):
    assert client.get("/api/reports/1/export").status_code == 401


def test_export_unknown_report_404(client, auth_headers):
    r = client.get("/api/reports/99999/export", headers=auth_headers)
    assert r.status_code == 404


def test_export_markdown_ok(client, auth_headers):
    rid = _gen_sync(client, auth_headers)
    r = client.get(f"/api/reports/{rid}/export", params={"fmt": "markdown"}, headers=auth_headers)
    assert r.status_code == 200
    assert "text/markdown" in r.headers.get("content-type", "")
    assert ".md" in r.headers.get("content-disposition", "")
    text = r.content.decode("utf-8")
    assert "GMV" in text and "296,087.91" in text and "订单转化率" in text


def test_export_docx_ok(client, auth_headers):
    rid = _gen_sync(client, auth_headers)
    r = client.get(f"/api/reports/{rid}/export", params={"fmt": "docx"}, headers=auth_headers)
    assert r.status_code == 200
    assert "application/vnd.openxmlformats" in r.headers.get("content-type", "")
    assert ".docx" in r.headers.get("content-disposition", "")
    blob = r.content
    assert blob[:2] == b"PK", "docx 应为 zip(PK) 魔数"
    text = _docx_text(blob)
    assert "GMV" in text, "docx 应包含 GMV 指标"
    assert "296,087.91" in text, "docx 金额应与预览一致"
    assert "2010-11-22" in text, "docx 应含逐日明细日期"
    assert "核心经营指标" in text or "经营" in text


def test_export_pending_rejected(client, auth_headers):
    # 手工造一条 pending 记录，导出应被拒(400)
    reg = client.post("/api/auth/register", json={"username": "bob", "password": "secret123"}).json()
    db = SessionLocal()
    try:
        user = db.query(User).filter_by(username="bob").first()
        rep = Report(user_id=user.id, instruction="昨日日报", status="pending")
        db.add(rep)
        db.commit()
        db.refresh(rep)
        rid = rep.id
    finally:
        db.close()
    headers = {"Authorization": f"Bearer {reg['access_token']}"}
    r = client.get(f"/api/reports/{rid}/export", headers=headers)
    assert r.status_code == 400
    assert "不可导出" in r.json()["detail"]


def test_export_cross_user_404(client, auth_headers):
    rid = _gen_sync(client, auth_headers)  # alice 的报告
    r2 = client.post("/api/auth/register", json={"username": "eve", "password": "secret123"})
    eve_headers = {"Authorization": f"Bearer {r2.json()['access_token']}"}
    r = client.get(f"/api/reports/{rid}/export", headers=eve_headers)
    assert r.status_code == 404  # 越权一律 404，不泄露存在性
